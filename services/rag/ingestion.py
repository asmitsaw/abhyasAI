import os
import hashlib
import io
from typing import Any, Dict, List, Optional, Tuple, Union
import chromadb
from chromadb.config import Settings
from services.document_service import extract_and_clean_document
from services.rag.chunking import StructureAwareChunker
from services.rag.embeddings import AbhyasEmbeddingFunction, get_collection_name
from services.rag.metadata import ChunkMetadata


CHROMA_DATA_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "data", "chroma")
os.makedirs(CHROMA_DATA_PATH, exist_ok=True)


class UniversalIngestionService:
    """
    Universal document ingestion pipeline supporting:
    PDF, TXT, Markdown, DOCX, PPTX, YouTube Transcripts, and Raw Pasted Text.
    Prevents duplicate ingestion via SHA-256 document hashing.
    """

    def __init__(self, chroma_path: Optional[str] = None):
        self.chroma_path = chroma_path or CHROMA_DATA_PATH
        self.client = chromadb.PersistentClient(path=self.chroma_path)
        self.embedding_fn = AbhyasEmbeddingFunction()
        self.chunker = StructureAwareChunker()

    def get_or_create_collection(self, collection_type: str):
        collection_name = get_collection_name(f"abhyas_{collection_type.lower()}")
        return self.client.get_or_create_collection(
            name=collection_name,
            embedding_function=self.embedding_fn,
            metadata={"hnsw:space": "cosine"}
        )

    def compute_hash(self, content: Union[str, bytes]) -> str:
        if isinstance(content, str):
            content_bytes = content.encode("utf-8", errors="ignore")
        else:
            content_bytes = content
        return hashlib.sha256(content_bytes).hexdigest()

    def is_document_ingested(self, collection, document_id: str) -> bool:
        try:
            results = collection.get(where={"document_id": document_id}, limit=1)
            return len(results.get("ids", [])) > 0
        except Exception:
            return False

    def extract_text(self, file_source: Any, filename: str) -> str:
        if not filename:
            if isinstance(file_source, str):
                return file_source.strip()
            return ""

        ext = os.path.splitext(filename)[1].lower()

        # Handle DOCX
        if ext in [".docx", ".doc"]:
            return self._extract_docx(file_source)

        # Handle PPTX
        elif ext in [".pptx", ".ppt"]:
            return self._extract_pptx(file_source)

        # Handle standard PDF, TXT, MD
        elif ext in [".pdf", ".txt", ".md", ".csv"]:
            return extract_and_clean_document(file_source, filename)

        # Fallback raw text
        else:
            try:
                return extract_and_clean_document(file_source, filename)
            except Exception:
                if hasattr(file_source, "read"):
                    data = file_source.read()
                    if isinstance(data, bytes):
                        return data.decode("utf-8", errors="ignore")
                    return str(data)
                return str(file_source)

    def _extract_docx(self, file_source: Any) -> str:
        try:
            import docx
            if isinstance(file_source, (str, os.PathLike)):
                doc = docx.Document(file_source)
            else:
                doc = docx.Document(file_source)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            return "\n\n".join(paragraphs)
        except Exception as error:
            print(f"DOCX extraction warning: {error}")
            return ""

    def _extract_pptx(self, file_source: Any) -> str:
        try:
            import pptx
            if isinstance(file_source, (str, os.PathLike)):
                prs = pptx.Presentation(file_source)
            else:
                prs = pptx.Presentation(file_source)
            slides_text = []
            for i, slide in enumerate(prs.slides, 1):
                slide_lines = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_lines.append(shape.text.strip())
                if slide_lines:
                    slides_text.append(f"--- SLIDE {i} ---\n" + "\n".join(slide_lines))
            return "\n\n".join(slides_text)
        except Exception as error:
            print(f"PPTX extraction warning: {error}")
            return ""

    def ingest_document(
        self,
        file_source: Any,
        filename: str,
        document_type: str,
        subject: str = "General",
        year: str = "",
        force_reindex: bool = False
    ) -> Dict[str, Any]:
        """
        Universal ingestion entry point.
        """
        raw_text = self.extract_text(file_source, filename)
        if not raw_text or not raw_text.strip():
            raise ValueError(f"Could not extract meaningful text from {filename or 'source'}.")

        doc_hash = self.compute_hash(raw_text)
        doc_type_clean = document_type.lower()
        target_collection_type = doc_type_clean if doc_type_clean in ["syllabus", "pyq", "lecture", "notes", "question", "answer"] else "documents"
        collection = self.get_or_create_collection(target_collection_type)

        # Check duplicate
        if not force_reindex and self.is_document_ingested(collection, doc_hash):
            existing = collection.get(where={"document_id": doc_hash})
            return {
                "document_id": doc_hash,
                "chunks_created": len(existing.get("ids", [])),
                "collection": collection.name,
                "status": "already_indexed",
                "message": "Document already indexed with identical hash. Vectors preserved."
            }

        # Structure-aware chunking
        chunk_tuples = self.chunker.chunk_document(
            text=raw_text,
            document_id=doc_hash,
            document_type=doc_type_clean,
            source_name=filename or "Uploaded Document",
            subject=subject,
            year=year
        )

        if not chunk_tuples:
            raise ValueError("Document produced 0 chunks after chunking.")

        ids = [meta.chunk_id for _, meta in chunk_tuples]
        documents = [txt for txt, _ in chunk_tuples]
        metadatas = [meta.to_chroma_dict() for _, meta in chunk_tuples]

        # Add to ChromaDB in batch
        collection.upsert(
            ids=ids,
            documents=documents,
            metadatas=metadatas
        )

        return {
            "document_id": doc_hash,
            "chunks_created": len(ids),
            "collection": collection.name,
            "status": "success",
            "subject": subject,
            "document_type": doc_type_clean
        }

    def ingest_youtube_transcript(
        self,
        video_id: str,
        transcript_text: str,
        subject: str = "General"
    ) -> Dict[str, Any]:
        source_name = f"YouTube Lecture ({video_id})"
        return self.ingest_document(
            file_source=transcript_text,
            filename=source_name,
            document_type="lecture",
            subject=subject
        )

    def ingest_raw_text(
        self,
        text: str,
        title: str,
        document_type: str,
        subject: str = "General",
        year: str = ""
    ) -> Dict[str, Any]:
        return self.ingest_document(
            file_source=text,
            filename=title,
            document_type=document_type,
            subject=subject,
            year=year
        )
