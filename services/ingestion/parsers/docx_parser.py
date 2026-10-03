import io
from typing import List
import docx
from services.document_service import clean_and_normalize_text
from services.ingestion.models import ExtractedChunk, NormalizedDocument


def parse_docx(file_bytes: bytes, filename: str) -> NormalizedDocument:
    """
    Extract headings and paragraphs from DOCX bytes.
    """
    doc = docx.Document(io.BytesIO(file_bytes))
    chunks: List[ExtractedChunk] = []
    text_blocks: List[str] = []
    current_heading = "General Notes"
    current_section_paras: List[str] = []
    page_approx = 1

    for para in doc.paragraphs:
        txt = clean_and_normalize_text(para.text)
        if not txt:
            continue

        # Check if heading style
        if para.style and "Heading" in para.style.name:
            if current_section_paras:
                sec_text = "\n".join(current_section_paras)
                text_blocks.append(f"### {current_heading}\n{sec_text}")
                chunks.append(ExtractedChunk(
                    text=sec_text,
                    page=page_approx,
                    section=current_heading,
                    heading=current_heading,
                    source_type="docx"
                ))
                current_section_paras = []
                page_approx = max(1, len(text_blocks) // 3 + 1)
            current_heading = txt
        else:
            current_section_paras.append(txt)

    if current_section_paras:
        sec_text = "\n".join(current_section_paras)
        text_blocks.append(f"### {current_heading}\n{sec_text}")
        chunks.append(ExtractedChunk(
            text=sec_text,
            page=page_approx,
            section=current_heading,
            heading=current_heading,
            source_type="docx"
        ))

    all_text = "\n\n".join(text_blocks)
    return NormalizedDocument(
        text=all_text,
        original_filename=filename,
        source_type="docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        file_extension=".docx",
        file_size=len(file_bytes),
        page_count=page_approx,
        chunks=chunks
    )
