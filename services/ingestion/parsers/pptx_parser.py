import io
from typing import List
import pptx
from services.document_service import clean_and_normalize_text
from services.ingestion.models import ExtractedChunk, NormalizedDocument


def parse_pptx(file_bytes: bytes, filename: str) -> NormalizedDocument:
    """
    Extract slide number, title, body content, and speaker notes from PowerPoint presentation.
    """
    prs = pptx.Presentation(io.BytesIO(file_bytes))
    chunks: List[ExtractedChunk] = []
    text_blocks: List[str] = []

    for i, slide in enumerate(prs.slides, 1):
        slide_title = f"Slide {i}"
        slide_texts = []

        for shape in slide.shapes:
            if shape.has_text_frame:
                txt = clean_and_normalize_text(shape.text_frame.text)
                if txt:
                    if shape == slide.shapes.title:
                        slide_title = txt
                    else:
                        slide_texts.append(txt)

        # Speaker notes if present
        notes_text = ""
        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                n_txt = clean_and_normalize_text(slide.notes_slide.notes_text_frame.text)
                if n_txt:
                    notes_text = f"\nSpeaker Notes: {n_txt}"
        except Exception:
            pass

        content_body = "\n".join(slide_texts) + notes_text
        if not content_body.strip():
            content_body = slide_title

        slide_block = f"--- SLIDE {i}: {slide_title} ---\n{content_body}"
        text_blocks.append(slide_block)

        chunks.append(ExtractedChunk(
            text=f"{slide_title}\n{content_body}".strip(),
            page=i,
            slide=i,
            section=f"Slide {i}",
            heading=slide_title,
            source_type="pptx"
        ))

    all_text = "\n\n".join(text_blocks)
    return NormalizedDocument(
        text=all_text,
        original_filename=filename,
        source_type="pptx",
        mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        file_extension=".pptx",
        file_size=len(file_bytes),
        page_count=len(prs.slides),
        chunks=chunks
    )
